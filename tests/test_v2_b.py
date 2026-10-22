"""Предпросмотр до подписи и порядок в итоговом Excel."""
from __future__ import annotations

from io import BytesIO

import openpyxl
from fastapi.testclient import TestClient

from chzkiz import chz_api, chz_mock, excel_out, kiz, sverka
from chzkiz.web.app import create_app
from tests.test_attacks_mod import _client, _code, _fin, _signin, _upload

INN = chz_mock.DEMO_OUR_INN
HEADER = ["КИЗ", "ИНН владельца", "Статус в отчете", "Статус в ЧЗ", "Новый статус", "Название товара",
          "Документ", "Номер документа", "Дата документа", "Сумма"]


def _seed(n: int, status: str, owner: str = INN) -> str:
    chz_mock.seed_cis(kiz.normalize_kiz(_code(n)), status, owner)
    return kiz.normalize_kiz(_code(n))


def _run(tg_id: int):
    vyvod = _seed(1, chz_api.STATUS_INTRODUCED)
    vozvrat = _seed(2, chz_api.STATUS_RETIRED)
    uzhe_vyveden = _seed(3, chz_api.STATUS_RETIRED)
    uzhe_v_oborote = _seed(4, chz_api.STATUS_INTRODUCED)
    chuzhoy = _seed(6, chz_api.STATUS_INTRODUCED, chz_mock.DEMO_OTHER_INN)
    c = _client(tg_id)
    f = _fin([("Продажа", _code(1), 5001, "500", "2026-10-01"),
              ("Возврат", _code(2), 5002, "600", "2026-10-02"),
              ("Продажа", _code(3), 5003, "700", "2026-10-03"),
              ("Возврат", _code(4), 5004, "800", "2026-10-04"),
              ("Продажа", _code(5), 5005, "900", "2026-10-05"),
              ("Продажа", _code(6), 5006, "950", "2026-10-06")])
    run_id, _ = _upload(c, [("report.xlsx", f)])
    _signin(c, run_id)
    return c, run_id, dict(vyvod=vyvod, vozvrat=vozvrat, uzhe_vyveden=uzhe_vyveden,
                           uzhe_v_oborote=uzhe_v_oborote, chuzhoy=chuzhoy, net=kiz.normalize_kiz(_code(5)))


def _book(content: bytes):
    return openpyxl.load_workbook(BytesIO(content))


def test_preview_otdaetsya_vladeltsu_chuzhomu_net():
    c, run_id, _ = _run(8101)
    r = c.get(f"/preview/{run_id}")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f"preview_{run_id[:8]}.xlsx" in r.headers["content-disposition"]
    assert _client(8102).get(f"/preview/{run_id}").status_code in (403, 404)
    assert TestClient(create_app()).get(f"/preview/{run_id}").status_code in (401, 403, 404)


def test_preview_listy_kolonki_i_oformlenie():
    c, run_id, _ = _run(8103)
    wb = _book(c.get(f"/preview/{run_id}").content)
    assert wb.sheetnames == ["Будут изменены", "Без изменений"]
    budut, bez = wb["Будут изменены"], wb["Без изменений"]
    assert [x.value for x in budut[1]] == HEADER
    assert [x.value for x in bez[1]] == HEADER[:4] + ["Объяснение"] + HEADER[5:]
    assert "Владелец" not in [x.value for x in budut[1]]
    for ws in (budut, bez):
        assert all(x.font.b for x in ws[1])
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == f"A1:J{ws.max_row}"
        assert ws.column_dimensions["A"].width >= max(len(str(x.value)) for x in ws["A"])
        assert ws.column_dimensions["E"].width >= max(len(str(x.value or "")) for x in ws["E"])


def test_preview_raskladka_na_mock_dannyh():
    c, run_id, k = _run(8104)
    wb = _book(c.get(f"/preview/{run_id}").content)
    budut = {r[0]: r for r in wb["Будут изменены"].iter_rows(min_row=2, values_only=True)}
    bez = {r[0]: r for r in wb["Без изменений"].iter_rows(min_row=2, values_only=True)}
    assert set(budut) == {k["vyvod"].replace("\x1d", " "), k["vozvrat"].replace("\x1d", " ")}
    assert len(bez) == 4
    kk = lambda key: k[key].replace("\x1d", " ")  # noqa: E731
    assert budut[kk("vyvod")][2:5] == ("Вывод", "В обороте", "Выбыл")
    assert budut[kk("vozvrat")][2:5] == ("Возврат", "Выбыл", "В обороте")
    assert budut[kk("vyvod")][1] == INN
    assert budut[kk("vyvod")][6] == "Заказ" and budut[kk("vyvod")][8] == "01.10.2026"
    assert budut[kk("vyvod")][9] == 500
    assert bez[kk("uzhe_vyveden")][4] == sverka.REASON_LABELS["already_output"]
    assert bez[kk("uzhe_v_oborote")][4] == sverka.REASON_LABELS["already_in_turnover"]
    assert bez[kk("net")][4] == sverka.REASON_LABELS["not_found"]
    assert bez[kk("chuzhoy")][4] == sverka.REASON_LABELS["wrong_owner"]
    assert bez[kk("uzhe_vyveden")][3] == "Выбыл" and bez[kk("uzhe_v_oborote")][3] == "В обороте"
    # номер документа это номер задания WB, как в файле изи чизи
    assert all(str(r[7]).isdigit() for r in list(budut.values()) + list(bez.values()))


def _row(i, cis, op, group, reason, status="INTRODUCED", **extra):
    r = {"id": i, "cis_raw": cis, "cis_norm": cis, "op": op, "group": group, "reason": reason,
         "chz_status": status, "owner_inn": INN, "product_cost": 159900, "date": "2026-09-29",
         "tovar": "Юбка карго"}
    r.update(extra)
    return r


def test_preview_nomer_zadaniya_est_i_bez_kolonok_ne_padaet():
    s_nomerom = _row(1, "0104600000000011AAA", "sale", sverka.GROUP_OUTPUT, "sale_in_turnover",
                     order_id=5550817538, istochnik="архив")
    bez_nomera = _row(2, "0104600000000011BBB", "sale", sverka.GROUP_OUTPUT, "sale_in_turnover")
    ws = _book(excel_out.build_preview([s_nomerom, bez_nomera]))["Будут изменены"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0][7] == "5550817538" and rows[0][5] == "Юбка карго" and rows[0][9] == 1599.0
    assert rows[1][7] in ("", None)
    assert _book(excel_out.build_preview([]))["Без изменений"].max_row == 1


def test_preview_prinyatyy_dokument_uhodit_v_bez_izmeneniy_i_opasnye_znacheniya():
    a = _row(1, "0104600000000011AAA", "sale", sverka.GROUP_OUTPUT, "sale_in_turnover", tovar="=1+1\x01")
    b = _row(2, "0104600000000011BBB", "sale", sverka.GROUP_OUTPUT, "sale_in_turnover")
    prn = _row(3, "0104600000000011CCC", "sale", sverka.GROUP_OK, "not_sold_yet")
    docs = [{"doc_type": chz_api.DOC_TYPE_RECEIPT, "status": chz_api.DOC_STATUS_CHECKED_OK,
             "codes_json": '["0104600000000011BBB"]'}]
    wb = _book(excel_out.build_preview([a, b, prn], docs))
    budut = list(wb["Будут изменены"].iter_rows(min_row=2, values_only=True))
    bez = {r[0]: r for r in wb["Без изменений"].iter_rows(min_row=2, values_only=True)}
    assert [r[0] for r in budut] == ["0104600000000011AAA"]
    assert budut[0][5] == "=1+1 " and wb["Будут изменены"]["F2"].data_type == "s"
    assert "принят" in bez["0104600000000011BBB"][4]
    assert bez["0104600000000011CCC"][2] == "-"


def test_otchet_pervye_listy_i_chto_delat_na_oshibkah():
    prichiny = ["not_found", "wrong_owner", "not_introduced", "bad_format", "duplicate", "wb_canceled"]
    rows = [_row(i, f"01046000000000{i}AAA", "sale", sverka.GROUP_ERROR, p, status=None)
            for i, p in enumerate(prichiny, 1)]
    rows.append(_row(9, "0104600000000099AAA", "sale", sverka.GROUP_ERROR, "weird_text"))
    rows.append(_row(10, "0104600000000010AAA", "sale", sverka.GROUP_OUTPUT, "sale_in_turnover"))
    wb = _book(excel_out.build_report(rows, []))
    assert wb.sheetnames[:3] == ["Изменены", "Без изменений", "Сводка"]
    assert wb.sheetnames[3:] == ["Выведено", "Возвращено", "Ошибки", "Все коды"]
    oshibki = list(wb["Ошибки"].iter_rows(values_only=True))
    assert oshibki[0][-1] == "Что делать"
    assert len(oshibki) == 8
    for r in oshibki[1:]:
        assert r[-1] and r[-1] != r[5], r
    assert "другому ИНН" in oshibki[2][-1]
    assert all(k in excel_out.REASON_ACTIONS for k in sverka.REASON_LABELS)


def test_shag3_inn_vladeltsa_i_knopka_predprosmotra():
    c, run_id, _ = _run(8105)
    html = c.get(f"/step3/{run_id}").text
    assert "<th>ИНН владельца</th>" in html
    assert f"<td>{INN}</td>" in html
    assert f'href="/preview/{run_id}"' in html and "Посмотреть что изменится" in html
    assert html.count('id="preview-link"') == 1
    assert html.index("Посмотреть что изменится") < html.index("<table")
