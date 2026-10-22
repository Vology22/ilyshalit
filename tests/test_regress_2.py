"""Атаки противника, круг 2: каждый тест тут должен падать на текущем коде. Не коммитить."""
from __future__ import annotations

import base64
import datetime
import io
import json
import re
import subprocess
import sys
import time
from pathlib import Path

import httpx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from chzkiz import chz_api, chz_mock, config, crypto, db, excel_in, sverka, wb_api
from chzkiz.web.app import create_app

APP_DIR = Path(__file__).resolve().parent.parent
# синтетические коды, не клиентские
GTIN = "04600000000017"
CODE = f"01{GTIN}21ATTACKSER0001"
CODE_RAW = f"{CODE}\\u001d91EE01\\u001d92" + "A" * 44
CODE2 = f"01{GTIN}21ATTACKSER0002"
INN = chz_mock.DEMO_OUR_INN
HDR = ("Код маркировки", "Тип документа", "Дата продажи")


def _mocksig(value: str) -> str:
    return base64.b64encode(f"MOCKSIG:{value}".encode()).decode("ascii")


def _login(c: TestClient, tg_id: int) -> None:
    with db.connect() as conn:
        user = db.get_or_create_user(conn, tg_id)
        token = db.create_login_token(conn, user["id"])
    r = c.post(f"/l/{token}", follow_redirects=True)
    assert r.status_code == 200


def _new_api_run(c: TestClient, tg_id: int) -> str:
    _login(c, tg_id)
    r = c.post("/step1", data={"source": "api"}, follow_redirects=False)
    assert r.status_code == 302
    return r.headers["location"].rsplit("/", 1)[-1]


def _excel_run(c: TestClient, tg_id: int, rows: list[tuple]) -> str:
    _login(c, tg_id)
    r = c.post("/step1", data={"source": "excel"},
               files={"file": ("r.xlsx", _xlsx(rows), "application/octet-stream")},
               follow_redirects=False)
    assert r.status_code == 302, r.text[:200]
    return r.headers["location"].rsplit("/", 1)[-1]


def _signin(c: TestClient, run_id: str, inn: str):
    p = c.post(f"/step2/{run_id}/auth-key").json()
    return c.post(
        f"/step2/{run_id}/signin",
        json={"uuid": p["uuid"], "signature": _mocksig(p["data"]), "inn": inn, "owner": "Тест"},
    )


def _signed_in_run(c: TestClient, tg_id: int, inn: str = INN) -> str:
    run_id = _new_api_run(c, tg_id)
    assert _signin(c, run_id, inn).status_code == 200
    return run_id


def _build(c: TestClient, run_id: str, doc_type: str) -> dict:
    r = c.post(f"/step4/{run_id}/build", json={"doc_type": doc_type, "paid": False})
    assert r.status_code == 200, r.text[:200]
    return r.json()


def _send(c: TestClient, run_id: str, d: dict, signature: str | None = None):
    return c.post(
        f"/step4/{run_id}/send",
        json={
            "doc_id": d["doc_id"],
            "chz_type": d["chz_type"],
            "product_document": d["product_document"],
            "signature": signature or _mocksig(d["product_document"]),
        },
    )


def _report_sheet_codes(c: TestClient, run_id: str, sheet: str) -> list[str]:
    r = c.get(f"/report/{run_id}")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content))[sheet]
    return [row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]


def _xlsx(rows: list[tuple]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _groups(run_id: str) -> list[str]:
    with db.connect() as conn:
        return [r["group"] for r in db.get_items(conn, run_id)]


# --- A. свертка продажа/возврат по дате мертва в живом потоке ---------------------

def test_live_flow_return_then_resale_goes_to_lk_return():
    rows = [HDR, (CODE_RAW, "Возврат", "2026-09-20"), (CODE_RAW, "Продажа", "2026-09-22")]
    info = {CODE: {"cis": CODE, "status": chz_api.STATUS_RETIRED, "ownerInn": INN}}
    unit = sverka.classify_items(excel_in.parse_report(_xlsx(rows)), info, INN)
    assert sverka.GROUP_RETURN not in [r["group"] for r in unit]  # на уровне функции с датой верно

    c = TestClient(create_app())
    chz_mock.seed_cis(CODE, chz_api.STATUS_RETIRED, INN)
    run_id = _excel_run(c, 9501, rows)
    assert _signin(c, run_id, INN).status_code == 200
    assert sverka.GROUP_RETURN not in _groups(run_id), (
        "последняя операция продажа 22.09, а живой поток шлет код на LK_RETURN: дата до сверки не доходит"
    )


def test_sale_return_resale_third_operation_dropped():
    rows = [HDR, (CODE_RAW, "Продажа", "2026-09-20"), (CODE_RAW, "Возврат", "2026-09-21"),
            (CODE_RAW, "Продажа", "2026-09-23")]
    items = excel_in.parse_report(_xlsx(rows))
    info = {CODE: {"cis": CODE, "status": chz_api.STATUS_RETIRED, "ownerInn": INN}}
    groups = [r["group"] for r in sverka.classify_items(items, info, INN)]
    assert sverka.GROUP_RETURN not in groups, (
        f"повторная продажа 23.09 выброшена дедупом (строк {len(items)}), код уходит на возврат: {groups}"
    )


def test_sale_return_pair_is_not_reported_as_error():
    rows = [HDR, (CODE_RAW, "Продажа", "2026-09-20"), (CODE_RAW, "Возврат", "2026-09-21")]
    items = excel_in.parse_report(_xlsx(rows))
    info = {CODE: {"cis": CODE, "status": chz_api.STATUS_RETIRED, "ownerInn": INN}}
    errs = [r["reason"] for r in sverka.classify_items(items, info, INN) if r["group"] == sverka.GROUP_ERROR]
    assert not errs, f"обычная пара продажа+возврат считается ошибкой: {errs}"


# --- B. документ завис в NEW: 409 навсегда -----------------------------------------

def test_build_without_send_locks_output_forever():
    c = TestClient(create_app())
    run_id = _signed_in_run(c, 9511)
    _build(c, run_id, "output")  # в браузере отменили ПИН-код или плагин упал, send не было
    r = c.post(f"/step4/{run_id}/build", json={"doc_type": "output"})
    locked = re.search(r'id="btn-output"[^>]*disabled', c.get(f"/step4/{run_id}").text) is not None
    assert r.status_code == 200, (
        f"повтор вывода после отмены подписи: {r.status_code}, кнопка вывода на шаге 4 заблокирована: {locked}"
    )


def test_send_rejected_by_chz_gives_500_and_locks():
    c = TestClient(create_app(), raise_server_exceptions=False)
    run_id = _signed_in_run(c, 9512)
    d = _build(c, run_id, "output")
    r = _send(c, run_id, d, signature=_mocksig("подпись не над этим документом"))
    r2 = c.post(f"/step4/{run_id}/build", json={"doc_type": "output"})
    assert r.status_code < 500 and r2.status_code == 200, (
        f"ЧЗ отклонил отправку 400: наружу {r.status_code}, повторная сборка {r2.status_code}"
    )


# --- C. разбор ответа doc/info ----------------------------------------------------

DOC_INFO_SHAPES = {
    "errors_null": {"status": chz_api.DOC_STATUS_CHECKED_NOT_OK, "errors": None},
    "list": [{"status": chz_api.DOC_STATUS_CHECKED_NOT_OK, "errors": ["код не найден"]}],
    "errors_str": {"status": chz_api.DOC_STATUS_CHECKED_NOT_OK, "errors": "код не найден"},
}


@pytest.mark.parametrize("shape", list(DOC_INFO_SHAPES))
def test_status_parses_doc_info_shapes(monkeypatch, shape):
    c = TestClient(create_app(), raise_server_exceptions=False)
    run_id = _signed_in_run(c, 9520)
    assert _send(c, run_id, _build(c, run_id, "output")).status_code == 200

    async def fake(self, doc_id):
        return DOC_INFO_SHAPES[shape]

    monkeypatch.setattr(chz_api.ChzClient, "doc_info", fake)
    st = c.get(f"/step4/{run_id}/status")
    assert st.status_code == 200, f"{shape}: опрос статуса {st.status_code}, документ навсегда IN_PROGRESS"
    doc = st.json()["documents"][0]
    assert doc["status"] == chz_api.DOC_STATUS_CHECKED_NOT_OK
    if shape != "errors_null":
        assert doc["errors"] == "код не найден", f"{shape}: причина в отчете {doc['errors']!r}"


# --- D. сводка в бот уходит один раз ----------------------------------------------

def test_bot_summary_after_return_done(monkeypatch):
    sent: list[str] = []
    real = httpx.AsyncClient

    def tg(req: httpx.Request) -> httpx.Response:
        sent.append(json.loads(req.content)["text"])
        return httpx.Response(200, json={"ok": True, "result": {}})

    def patched(*args, **kwargs):
        if kwargs.get("transport") is None:
            kwargs["transport"] = httpx.MockTransport(tg)
        return real(*args, **kwargs)

    monkeypatch.setattr(httpx, "AsyncClient", patched)
    monkeypatch.setattr(config, "BOT_TOKEN", "42:TEST")
    c = TestClient(create_app())
    run_id = _signed_in_run(c, 9530)
    _send(c, run_id, _build(c, run_id, "output"))
    assert c.get(f"/step4/{run_id}/status").json()["done"] is True
    _send(c, run_id, _build(c, run_id, "return"))
    assert c.get(f"/step4/{run_id}/status").json()["done"] is True
    assert _report_sheet_codes(c, run_id, "Возвращено"), "отчет возврат видит"
    assert any("Возвращено: 1" in t for t in sent), (
        f"возврат принят ЧЗ, а в бот ушло сообщений: {len(sent)}, последнее: {sent[-1:]}"
    )


# --- F. период на шаге 1 игнорируется ------------------------------------------------

def test_step1_api_period_from_form_is_ignored(monkeypatch):
    seen: list[dict] = []
    real = httpx.AsyncClient

    def wb(req: httpx.Request) -> httpx.Response:
        seen.append(dict(req.url.params))
        return httpx.Response(200, json={"orders": [], "next": 0})

    def patched(*args, **kwargs):
        if kwargs.get("transport") is None:
            kwargs["transport"] = httpx.MockTransport(wb)
        return real(*args, **kwargs)

    c = TestClient(create_app())
    _login(c, 9540)
    with db.connect() as conn:
        user = db.get_or_create_user(conn, 9540)
        db.set_wb_key(conn, user["id"], crypto.encrypt_str("test-key-not-real"))
    monkeypatch.setattr(config, "CHZ_MODE", "real")
    monkeypatch.setattr(httpx, "AsyncClient", patched)
    day = 86400
    d_from = datetime.date.fromtimestamp(time.time() - 20 * day)
    d_to = datetime.date.fromtimestamp(time.time() - 15 * day)
    c.post("/step1", data={"source": "api", "date_from": d_from.isoformat(), "date_to": d_to.isoformat()},
           follow_redirects=False)
    assert seen, "в WB не ходили"
    d_start = d_to - datetime.timedelta(days=29)
    want = time.mktime(d_start.timetuple()) - 60 * 86400  # задания берем с запасом в 60 дней назад
    got = int(seen[0]["dateFrom"])
    assert abs(got - want) <= day and "dateTo" in seen[0], (
        f"выбран период {d_from}..{d_to}, в WB ушло dateFrom={datetime.date.fromtimestamp(got)}, "
        f"dateTo={'есть' if 'dateTo' in seen[0] else 'нет'}"
    )


# --- G. 429 на meta WB молча теряет КИЗ ----------------------------------------------

async def test_wb_meta_429_silently_drops_kiz(monkeypatch):
    def handler(request: httpx.Request) -> httpx.Response:
        p = request.url.path
        if p == "/api/v3/orders":
            return httpx.Response(200, json={"next": 0, "orders": [{"id": 101, "price": 1000}, {"id": 102, "price": 1000}]})
        if p == "/api/v3/orders/status":
            return httpx.Response(200, json={"orders": [{"id": 101, "wbStatus": "sold"}, {"id": 102, "wbStatus": "sold"}]})
        if p == "/api/marketplace/v3/orders/meta":
            if json.loads(request.content)["orders"] == [101]:
                return httpx.Response(429, json={"title": "too many requests"})
            return httpx.Response(200, json={"orders": [{"id": 102, "meta": {"sgtin": {"value": [CODE2]}}}]})
        return httpx.Response(404, json={})

    real = httpx.AsyncClient

    def patched(*args, **kwargs):
        kwargs["transport"] = httpx.MockTransport(handler)
        return real(*args, **kwargs)

    monkeypatch.setattr(httpx, "AsyncClient", patched)
    monkeypatch.setattr(wb_api, "META_BATCH", 1)
    try:
        items = await wb_api.WbClient("test-key-not-real").fetch_items()
    except httpx.HTTPError:
        return  # громкий отказ допустим
    assert len(items) == 2, f"WB ответил 429 на один заказ, КИЗ пропал без ошибки: {len(items)} из 2"


# --- H. товарная группа пользователя не используется ----------------------------------

def test_send_ignores_user_product_group(monkeypatch):
    seen: list[httpx.Request] = []
    real = httpx.AsyncClient

    async def rec(req: httpx.Request) -> None:
        seen.append(req)

    def patched(*args, **kwargs):
        kwargs["event_hooks"] = {"request": [rec]}
        return real(*args, **kwargs)

    monkeypatch.setattr(httpx, "AsyncClient", patched)
    c = TestClient(create_app())
    run_id = _signed_in_run(c, 9560)
    with db.connect() as conn:
        conn.execute("UPDATE users SET pg = 'shoes' WHERE tg_id = ?", (9560,))
        conn.commit()
    _send(c, run_id, _build(c, run_id, "output"))
    creates = [r for r in seen if r.url.path.endswith(chz_api.DOC_CREATE_PATH)]
    assert creates, "документ не отправлялся"
    assert creates[0].url.params.get("pg") == "shoes", f"у пользователя shoes, в ЧЗ ушло pg={creates[0].url.params.get('pg')}"


# --- I. ИНН из SubjectName ------------------------------------------------------------

def _js_extract_inn(owner: str) -> str:
    js = (APP_DIR / "chzkiz" / "web" / "static" / "app.js").read_text(encoding="utf-8")
    body = js.split("function chzExtractInn(owner) {", 1)[1].split("\n}", 1)[0]
    for rx in re.findall(r"/(.+?)/i\.exec\(owner", body):
        m = re.search(rx, owner, re.IGNORECASE)
        if m:
            return m.group(1)
    return ""


@pytest.mark.parametrize(
    "subject,want",
    [
        ("CN=ООО Тест, O=ООО Тест, INNLE=7700000000, OGRN=1027700000000, SN=Иванов, INN=770000000001", "7700000000"),
        ("CN=ООО Тест, OID.1.2.643.100.4=7700000000, OID.1.2.643.3.131.1.1=770000000001", "7700000000"),
        ("CN=ИП Тест, SN=Тест, OID.1.2.643.3.131.1.1=770000000001, OGRNIP=300000000000000", "770000000001"),
    ],
    ids=["INNLE", "OID_UL", "OID_IP"],
)
def test_inn_extract_subject_variants(subject, want):
    got = _js_extract_inn(subject)
    assert got == want, f"из SubjectName взят ИНН {got!r} вместо {want}"


# --- J. цена с разделителем тысяч -----------------------------------------------------

@pytest.mark.parametrize("cell", ["1 234,56", "1 234,56"], ids=["space", "nbsp"])
def test_price_with_thousand_separator_becomes_zero(cell):
    rows = [("Код маркировки", "Тип документа", "Цена розничная с учетом согласованной скидки"),
            (CODE_RAW, "Продажа", cell)]
    items = excel_in.parse_report(_xlsx(rows))
    assert items[0]["product_cost"] == 123456, f"цена {cell!r} стала {items[0]['product_cost']} копеек"


# --- K. бот: ключ вне состояния WaitKey --------------------------------------------------

def _bot_env():
    from aiogram import Bot
    from aiogram.client.session.base import BaseSession
    from aiogram.methods import SendMessage
    from aiogram.types import Chat, Message

    class FakeSession(BaseSession):
        def __init__(self):
            super().__init__()
            self.calls = []

        async def close(self):
            return None

        async def make_request(self, bot, method, timeout=None):
            self.calls.append(method)
            if isinstance(method, SendMessage):
                return Message(message_id=900 + len(self.calls), date=datetime.datetime.now(),
                               chat=Chat(id=method.chat_id, type="private"), text=method.text)
            return True

        async def stream_content(self, url, headers=None, timeout=30, chunk_size=65536, raise_for_status=True):
            if False:
                yield b""

    from chzkiz.bot import main as bot_main

    session = FakeSession()
    return bot_main, Bot(token="42:TEST", session=session), session


def _upd(uid: int, text: str, n: int):
    from aiogram.types import Chat, Message, Update, User

    return Update(update_id=n, message=Message(
        message_id=n, date=datetime.datetime.now(), chat=Chat(id=uid, type="private"),
        from_user=User(id=uid, is_bot=False, first_name="T"), text=text))


async def test_bot_key_after_restart_stays_in_chat(monkeypatch):
    from aiogram.methods import DeleteMessage

    bot_main, bot, session = _bot_env()

    async def ok(raw_key):
        return True

    monkeypatch.setattr(bot_main, "_check_wb_key", ok)
    dp1 = bot_main.build_dispatcher()
    await dp1.feed_update(bot, _upd(9580, "/start", 1))
    dp2 = bot_main.build_dispatcher()  # перезапуск службы между /start и ключом: MemoryStorage пуст
    await dp2.feed_update(bot, _upd(9580, "eyJhbGciOiJIUzI1NiJ9.testkeynotreal0123456789.signature", 2))
    deleted = [m for m in session.calls if isinstance(m, DeleteMessage)]
    assert deleted, f"ключ WB остался в чате без ответа, вызовов бота после /start: {len(session.calls) - 1}"


# --- L. бомба по ширине строк данных при узком заголовке ------------------------------------

def test_wide_data_rows_bomb_blows_memory(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Код маркировки"
    ws["B1"] = "Тип документа"
    for i in range(2, 12002):
        ws.cell(row=i, column=16384, value=1)
    path = tmp_path / "wide.xlsx"
    wb.save(path)
    size = path.stat().st_size
    assert size < 200_000
    script = (
        "import resource,sys\n"
        "lim=1024*1024*1024\n"
        "resource.setrlimit(resource.RLIMIT_AS,(lim,lim))\n"
        "from chzkiz import excel_in\n"
        "data=open(sys.argv[1],'rb').read()\n"
        "try:\n"
        "    excel_in.parse_report(data)\n"
        "except excel_in.ExcelColumnError:\n"
        "    pass\n"
        "print('OK')\n"
    )
    p = subprocess.run([sys.executable, "-c", script, str(path)],
                       cwd=str(APP_DIR), capture_output=True, text=True, timeout=120)
    assert p.returncode == 0 and "OK" in p.stdout, (
        f"файл {size // 1024} КБ, заголовок 2 колонки: разбор не уложился в 1 ГБ: {p.stderr.strip()[-160:]}"
    )
