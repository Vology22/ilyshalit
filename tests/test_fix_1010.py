"""Пакет правок 10.10: архив без ключа WB, возврат из финотчета, шаг 4, тексты, бот, отчет, страницы ошибок."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from aiogram.methods import DeleteMessage, SendMessage
from aiogram.types import Chat, Dice, Message, Update, User

from chzkiz import chz_api, chz_errors, config, db, excel_in, excel_out, kiz, sverka, wb_api
from chzkiz.web import app as app_mod
from tests.test_attacks_mod import _client, _code, _fin, _rows, _save, _signin, _upload
from tests.test_regress_1 import _bot_env, _upd

PKG = Path(config.__file__).resolve().parent
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _arch_real(rows) -> bytes:
    """Как настоящий archive.xlsx WB, без покупателей. rows: (задание, код, цена, статус, тип операции, дата)."""
    wb = openpyxl.Workbook()
    tasks = wb.active
    tasks.title = "Сборочные задания"
    tasks.append(["№ задания", "Артикул", "Статус задания"])
    sheet = wb.create_sheet("КИЗ")
    sheet.append(["№ задания", "Стикер", "КИЗ", "Номер чека", "Стоимость", "Валюта",
                  "Номер фискального накопителя", "Дата", "Тип операции", "Признак продажи юрлицу"])
    for order, code, price, status, op, day in rows:
        tasks.append([order, "A", status])
        sheet.append([order, "1:2", code, "", price, "RUB", "", day, op, "Нет"])
    return _save(wb)


@pytest.fixture
def real(monkeypatch):
    monkeypatch.setattr(app_mod.config, "is_mock", lambda: False)


class _FakeWb:
    table = {1: "sold", 2: "canceled_by_client", 3: "ready_for_pickup"}

    def __init__(self, key):
        pass

    async def get_orders_status(self, ids):
        return {i: {"wbStatus": self.table[i]} for i in ids if i in self.table}


def _with_key(monkeypatch, tg_id: int):
    monkeypatch.setattr(app_mod.wb_api, "WbClient", _FakeWb)
    monkeypatch.setattr(app_mod, "_decrypt", lambda v: v)
    with db.connect() as conn:
        db.set_wb_key(conn, db.get_or_create_user(conn, tg_id)["id"], "k")
    return _client(tg_id)


# --- пункт 1: архив без ключа WB ---------------------------------------------------------------

def test_parse_arhiva_beret_status_operaciyu_i_datu_bez_pokupatelya():
    f = _arch_real([(1, _code(1), "100", "Продано", "Продажа", "11:51:00 07.10.2026"),
                    (2, _code(2), "100", "Отказ покупателем", "-", None),
                    (3, _code(3), "100", None, "Возврат", None)])
    by = {it["order_id"]: it for it in excel_in.parse_report(f)}
    assert (by[1]["arh_status"], by[1]["arh_op"], by[1]["arh_date"]) == ("Продано", "Продажа", "2026-10-07")
    assert (by[2]["arh_status"], by[2]["arh_op"], by[2]["arh_date"]) == ("Отказ покупателем", "", "")
    assert (by[3]["arh_status"], by[3]["arh_op"]) == ("", "Возврат")
    assert not any("покупател" in k or "телефон" in k for it in by.values() for k in it)


def test_arhiv_bez_klyucha_kladet_statusy_iz_faila(real):
    sts = ["Продано", "Прибыл на ПВЗ", "Отказ покупателем", "Покупатель отменил заказ", "Брак", "Отсортировано"]
    f = _arch_real([(10 + i, _code(i), "100", s, "-", None) for i, s in enumerate(sts)])
    _, rows = _upload(_client(9601), [("archive.xlsx", f)])
    assert [(r["op"], r["reason"]) for r in rows] == [
        ("sale", ""), ("sale", "not_sold_yet"), ("return", "otkaz"),
        ("return", "otkaz"), ("sale", "not_sold_yet"), ("sale", "not_sold_yet")]


def test_arhiv_bez_statusa_beret_tip_operacii(real):
    f = _arch_real([(20, _code(20), "100", None, "Продажа", None),
                    (21, _code(21), "100", None, "Возврат", None),
                    (22, _code(22), "100", None, "-", None)])
    _, rows = _upload(_client(9602), [("archive.xlsx", f)])
    assert [(r["op"], r["reason"]) for r in rows] == [
        ("sale", ""), ("return", "otkaz"), ("sale", "not_sold_yet")]


def test_s_klyuchom_wb_raskladka_prezhnyaya(real, monkeypatch):
    c = _with_key(monkeypatch, 9603)
    f = _arch_real([(1, _code(31), "100", "Прибыл на ПВЗ", "-", None),
                    (2, _code(32), "100", "Продано", "Продажа", None),
                    (3, _code(33), "100", "Продано", "Продажа", None)])
    _, rows = _upload(c, [("archive.xlsx", f)])
    assert [(r["op"], r["reason"]) for r in rows] == [
        ("sale", ""), ("return", "otkaz"), ("sale", "not_sold_yet")]


# --- пункт 2: архив вместе с финотчетом ----------------------------------------------------------

def test_vozvrat_iz_finotcheta_delaet_stroku_arhiva_vozvratom(real):
    f = _fin([("Продажа", _code(40), 1, "500", "2026-10-01"), ("Возврат", _code(40), 1, "500", "2026-10-03")])
    a = _arch_real([(1, _code(40), "", "Прибыл на ПВЗ", "-", None)])
    _, rows = _upload(_client(9604), [("report.xlsx", f), ("archive.xlsx", a)])
    assert len(rows) == 1
    assert (rows[0]["op"], rows[0]["reason"], rows[0]["date"]) == ("return", "vozvrat_fin", "2026-10-03")
    assert rows[0]["product_cost"] == 50000


def test_prodazha_iz_finotcheta_snimaet_not_sold_yet(real):
    f = _fin([("Продажа", _code(41), 1, "500", "2026-10-01")])
    a = _arch_real([(1, _code(41), "", "Прибыл на ПВЗ", "-", None)])
    _, rows = _upload(_client(9605), [("report.xlsx", f), ("archive.xlsx", a)])
    assert len(rows) == 1
    assert (rows[0]["op"], rows[0]["reason"], rows[0]["date"]) == ("sale", "", "2026-10-01")


def test_vozvrat_finotcheta_s_klyuchom_wb_i_po_zadaniyu_bez_koda(real, monkeypatch):
    c = _with_key(monkeypatch, 9606)
    f = _fin([("Возврат", "", 1, "500", "2026-10-05")])
    a = _arch_real([(1, _code(42), "100", "Продано", "Продажа", "09:00:00 02.10.2026")])
    _, rows = _upload(c, [("report.xlsx", f), ("archive.xlsx", a)])
    assert [(r["op"], r["reason"], r["date"]) for r in rows] == [("return", "vozvrat_fin", "2026-10-05")]


def test_kody_finotcheta_bez_arhiva_ne_dobavlyayutsya_i_data_iz_arhiva(real):
    f = _fin([("Продажа", _code(43), 9, "500", "2026-10-01")])
    a = _arch_real([(1, _code(44), "100", "Продано", "Продажа", "11:51:00 07.10.2026")])
    _, rows = _upload(_client(9607), [("report.xlsx", f), ("archive.xlsx", a)])
    assert [(kiz.normalize_kiz(r["cis_raw"]), r["date"]) for r in rows] == [(kiz.normalize_kiz(_code(44)), "2026-10-07")]


@pytest.mark.parametrize("chz_status,group,reason", [
    (chz_api.STATUS_RETIRED, sverka.GROUP_RETURN, "vozvrat_fin_vyveden"),
    (chz_api.STATUS_INTRODUCED, sverka.GROUP_OK, "vozvrat_fin_v_oborote"),
])
def test_vozvrat_fin_v_gruppah_sverki(chz_status, group, reason):
    code = wb_api._fake_cis("781002")
    it = {"cis_raw": code, "op": wb_api.OP_RETURN, "reason": "vozvrat_fin", "product_cost": 100, "date": "2026-10-03"}
    info = {kiz.normalize_kiz(code): {"status": chz_status, "ownerInn": "1"}}
    r = sverka.classify_items([it], info, "1")[0]
    assert (r["group"], r["reason"]) == (group, reason)
    assert "возврат в финотчете" in sverka.REASON_LABELS[reason]
    assert sverka.REASON_LABELS["vozvrat_fin"] == "покупатель вернул товар, возврат в финотчете"


# --- пункт 3: шаг 4 после обновления -------------------------------------------------------------

def test_app_js_podklyuchen_s_novoy_versiey_vezde():
    for name in ("step2.html", "step4.html"):
        page = (PKG / "web" / "templates" / name).read_text(encoding="utf-8")
        assert "/static/app.js?v=1310a" in page and "app.js?v=1010\"" not in page, name


def test_shag4_pervyy_opros_srazu_esli_est_dokument():
    js = (PKG / "web" / "static" / "app.js").read_text(encoding="utf-8")
    assert "if (hasDocs) { chzPollStatus(runId, statusBlock, true); }" in js
    assert "if (immediate) { tick(); }" in js
    c = _client(9610)
    run_id, _ = _upload(c, [("archive.xlsx", _arch_real([(1, _code(50), "100", "Продано", "Продажа", None)]))])
    _signin(c, run_id)
    assert c.get(f"/step3/{run_id}").status_code == 200
    page = c.get(f"/step4/{run_id}").text
    assert "v=1310a" in page and ", false);</script>" in page


# --- пункт 4: срок входа ---------------------------------------------------------------------------

def test_vhod_zhivet_30_dnej_vo_vseh_tekstah():
    assert config.SESSION_TTL_HOURS == 720
    for path in PKG.rglob("*"):
        if path.suffix in (".py", ".html", ".js") and not path.name.startswith("._"):
            assert "12 час" not in path.read_text(encoding="utf-8"), path
    c = _client(9611)
    assert "30 дней" in c.get("/instrukciya").text
    from chzkiz.bot import main as bot_main
    assert any("30 дней" in body for _, body in bot_main.INSTRUKCIYA_STEPS)


# --- пункт 5 -----------------------------------------------------------------------------------

def test_instrukciya_bota_pro_fias_bez_prosby_vvesti():
    from chzkiz.bot import main as bot_main
    text = dict(bot_main.INSTRUKCIYA_STEPS)["Вывод и возврат"]
    assert "потом возврат. Код ФИАС склада панель подставит сама" in text
    assert "укажите код ФИАС" not in text


def test_pyas_klikov_i_placeholder_fias():
    from chzkiz.bot import screens
    assert "Осталась пара кликов" in screens.go_screen("http://x")[0]
    page = (PKG / "web" / "templates" / "step4.html").read_text(encoding="utf-8")
    assert 'placeholder="36 знаков, например 0c5b2444-70a0-..."' in page
    assert "В документе цена из отчета, дата вывода сегодняшняя." in page


def test_shag3_podskazka_pro_listanie_tolko_na_uzkom_ekrane():
    page = (PKG / "web" / "templates" / "step3.html").read_text(encoding="utf-8")
    css = (PKG / "web" / "static" / "style.css").read_text(encoding="utf-8")
    assert "Таблицу можно листать вбок" in page
    assert ".scroll-hint { display: none;" in css and "@media (max-width: 700px)" in css


def test_chz_errors_date_tolko_celym_slovom():
    assert "датой" not in chz_errors.explain("validation failed")
    assert "датой" not in chz_errors.explain("update error")
    assert "датой" in chz_errors.explain("invalid date")
    assert "датой" in chz_errors.explain("Неверная дата документа")


def test_pohozhe_li_na_klyuch_wb():
    from chzkiz.bot.main import _looks_like_wb_key as ok
    jwt = "eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxMjM0.c2lnbmF0dXJlX3Rlc3Q"
    assert ok(jwt) and ok(f"  {jwt}\n")
    assert not ok("https://example.com/very/long/path/to/some/page")
    assert not ok("eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxMjM0")
    assert not ok("eyJhbGciOiJ.IUzI1NiJ9.eyJzdWIiOiIx.MjM0c2lnbg")
    assert not ok("eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiIxMjM0.c2lnbmF0dXJl+test")
    assert not ok("abcdefghij.klmnopqrstuvw.xyzabcdefgh")
    assert not ok("eyJ.a.b")


async def test_bot_otvechaet_podskazkoj_na_neponyatnoe():
    bot_main, bot, session = _bot_env()
    dp = bot_main.build_dispatcher()
    await dp.feed_update(bot, _upd(9620, "/nepoimu", 1))
    await dp.feed_update(bot, _upd(9620, "просто текст", 2))
    await dp.feed_update(bot, _upd(9620, "https://example.com/aaaaaaaaaaaaaaaaaaaaaaaa", 3))
    await dp.feed_update(bot, Update(update_id=4, message=Message(
        message_id=4, date=_upd(1, "", 1).message.date, chat=Chat(id=9620, type="private"),
        from_user=User(id=9620, is_bot=False, first_name="T"), dice=Dice(emoji="🎲", value=3))))
    texts = [m.text for m in session.calls if isinstance(m, SendMessage)]
    assert texts == [bot_main.HINT_TEXT] * 4, texts
    assert not [m for m in session.calls if isinstance(m, DeleteMessage)]
    await dp.feed_update(bot, _upd(9620, "/out", 5))
    assert [m.text for m in session.calls if isinstance(m, SendMessage)][-1] != bot_main.HINT_TEXT


def test_otchet_statusy_chz_po_russki_i_chto_eto_znachit():
    code = _code(60)
    base = {"id": 1, "cis_raw": code, "cis_norm": code, "op": "sale", "owner_inn": "1", "group": "error",
            "product_cost": 100, "date": "2026-10-01"}
    rows = [dict(base, chz_status="INTRODUCED", reason="duplicate"),
            dict(base, id=2, chz_status="RETIRED", reason="weird_reason"),
            dict(base, id=3, chz_status="NEW_STATUS", reason="duplicate")]
    wb = openpyxl.load_workbook(BytesIO(excel_out.build_report(rows, [])))
    err = list(wb["Ошибки"].iter_rows(values_only=True))
    assert err[0][6] == "Что делать"
    assert [r[3] for r in err[1:]] == ["в обороте", "выведен из оборота", "NEW_STATUS"]
    assert err[1][6] == excel_out.REASON_ACTIONS["duplicate"]
    assert err[2][6] == chz_errors.explain("weird_reason")
    assert excel_out.CHZ_STATUS_RU["DISAGGREGATION"] == "расформирован"
    assert excel_out.CHZ_STATUS_RU["WITHDRAWN"] == "выведен из оборота"


def test_stranicy_oshibok_v_obschem_shablone_a_json_ne_tronut():
    c = _client(9630)
    html = {"accept": "text/html"}
    r = c.get("/step2/net-takogo", headers=html)
    assert r.status_code == 404 and "Запуск не найден." in r.text
    assert "Вернуться к загрузке" in r.text and 'href="/step1"' in r.text and 'class="topbar"' in r.text
    r = c.get("/net-takoj-stranicy", headers=html)
    assert r.status_code == 404 and "Страница не найдена." in r.text and "topbar" in r.text
    r = c.post("/step1", data={"source": "excel"})
    assert r.status_code == 400 and "Файл не выбран." in r.text and "Вернуться к загрузке" in r.text
    r = c.get("/step2/net-takogo")
    assert r.status_code == 404 and r.json() == {"detail": "Запуск не найден."}
    r = c.post("/step4/net-takogo/build", json={"doc_type": "output"})
    assert r.status_code == 404 and r.json()["detail"] == "Запуск не найден."
    assert "r.json()" in (PKG / "web" / "static" / "app.js").read_text(encoding="utf-8")
