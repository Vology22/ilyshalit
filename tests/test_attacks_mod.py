"""Атаки на правку панели: два файла WB, место деятельности, разбивка на шаге 3."""
from __future__ import annotations

import base64
import json
from io import BytesIO

import httpx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from chzkiz import chz_api, chz_mock, db, kiz, wb_api
from chzkiz.web import app as app_mod
from chzkiz.web.app import create_app

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GS = "\\u001d"


def _code(n: int) -> str:
    return f"0104600000000{n:03d}21ABCDE{n:04d}{GS}91EE11{GS}92ZZZZ{n:04d}"


def _save(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fin(rows) -> bytes:
    """rows: (операция, код, номер задания, цена, дата)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Тип документа", "Обоснование для оплаты", "Код маркировки", "Номер сборочного задания",
               "Вайлдберриз реализовал Товар (Пр)", "Дата продажи"])
    for op, code, order, price, day in rows:
        ws.append([op, op, code, order, price, day])
    return _save(wb)


def _arch(rows) -> bytes:
    """rows: (номер задания, код, стоимость)."""
    wb = openpyxl.Workbook()
    tasks = wb.active
    tasks.title = "Сборочные задания"
    tasks.append(["№ задания", "Статус задания", "Артикул"])
    sheet = wb.create_sheet("КИЗ")
    sheet.append(["№ задания", "Стикер", "КИЗ", "Номер чека", "Стоимость", "Валюта",
                  "Номер фискального накопителя", "Дата", "Тип операции", "Признак продажи юрлицу"])
    for order, code, price in rows:
        tasks.append([order, "Доставлен", "A"])
        sheet.append([order, "1:2", code, "", price, "RUB", "", None, "-", "Нет"])
    return _save(wb)


def _client(tg_id: int) -> TestClient:
    with db.connect() as conn:
        user = db.get_or_create_user(conn, tg_id)
        token = db.create_login_token(conn, user["id"])
    c = TestClient(create_app())
    c.post(f"/l/{token}", follow_redirects=True)
    return c


def _upload(c: TestClient, files) -> tuple[str, list[dict]]:
    r = c.post("/step1", data={"source": "excel"},
               files=[("file", (n, b, XLSX)) for n, b in files], follow_redirects=False)
    assert r.status_code == 302, r.text[:300]
    run_id = r.headers["location"].rsplit("/", 1)[-1]
    return run_id, _rows(run_id)


def _rows(run_id: str) -> list[dict]:
    with db.connect() as conn:
        return [dict(x) for x in db.get_items(conn, run_id)]


def _signin(c: TestClient, run_id: str, inn: str = chz_mock.DEMO_OUR_INN) -> None:
    p = c.post(f"/step2/{run_id}/auth-key").json()
    sig = base64.b64encode(f"MOCKSIG:{p['data']}".encode()).decode("ascii")
    r = c.post(f"/step2/{run_id}/signin",
               json={"uuid": p["uuid"], "signature": sig, "inn": inn, "owner": "Тест"})
    assert r.status_code == 200, r.text


def test_prodazha_vozvrat_prodazha_odnogo_koda_ne_teryaet_vtoruyu_prodazhu():
    """Дата вылетела из ключа схлопывания: две продажи одного кода в разные дни слипаются,
    сверка видит продажа-возврат и считает код возвращенным, хотя последней была продажа."""
    c = _client(9401)
    f = _fin([("Продажа", _code(7), 5001, "500", "2026-10-01"),
              ("Возврат", _code(7), 5001, "500", "2026-10-02"),
              ("Продажа", _code(7), 5001, "700", "2026-10-03")])
    run_id, rows = _upload(c, [("report.xlsx", f)])
    assert len(rows) == 3, [(r["op"], r["date"]) for r in rows]
    _signin(c, run_id)
    rows = _rows(run_id)
    assert [r["op"] for r in rows] == [wb_api.OP_SALE], [(r["op"], r["date"]) for r in rows]


def test_fbs_stroka_otcheta_bez_koda_ne_pobezhdaet_arhiv(monkeypatch):
    """В финотчете у FBS код пустой, его потом дает WB. Схлопывание идет до подстановки кода,
    поэтому пара архив+отчет по одному заданию доживает до сверки двумя строками, и в
    сверке побеждает строка отчета (у нее есть дата), а не архива."""
    async def fake_fill(wb_key_enc, items):
        return [{**it, "cis_raw": _code(1), "need_meta": False} if it.get("need_meta") else it
                for it in items], 0

    monkeypatch.setattr(app_mod, "_fill_fbs_codes", fake_fill)
    c = _client(9402)
    run_id, rows = _upload(c, [
        ("report.xlsx", _fin([("Продажа", "", 1000, "500", "2026-10-01")])),
        ("archive.xlsx", _arch([(1000, _code(1), "1 234,50")])),
    ])
    assert len({kiz.normalize_kiz(r["cis_raw"]) for r in rows}) == len(rows) == 1, \
        [(r["istochnik"], r["date"]) for r in rows]
    _signin(c, run_id)
    assert [r["istochnik"] for r in _rows(run_id)] == ["архив"]


def test_arhiv_bez_ceny_ne_zatiraet_cenu_iz_otcheta():
    """Строка архива главная, даже если у нее цены нет, а у такой же в отчете цена есть.
    Вывод без цены панель не собирает, так что код встает колом при готовой цене в отчете."""
    c = _client(9403)
    run_id, rows = _upload(c, [
        ("report.xlsx", _fin([("Продажа", _code(2), 1002, "500", "2026-10-01")])),
        ("archive.xlsx", _arch([(1002, _code(2), "")])),
    ])
    assert len(rows) == 1
    assert rows[0]["product_cost"] == 50000, rows[0]["product_cost"]


def _chz_with(spisok, kartochki) -> chz_api.ChzClient:
    """ЧЗ-заглушка: doc/list отдает spisok, doc/{number}/info берет из kartochki (код, тело)."""
    def handler(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        if path == chz_api.DOC_LIST_PATH:
            return httpx.Response(200, json={"results": spisok})
        number = path.split("/")[-2]
        status, body = kartochki[number]
        if isinstance(body, bytes):
            return httpx.Response(status, content=body)
        return httpx.Response(status, json=body)

    return chz_api.ChzClient("http://chz.test", token="t", transport=httpx.MockTransport(handler))


async def test_bitaya_kartochka_ne_dolzhna_obryvat_poisk_po_ostalnym_dokumentam():
    """Шлюз вернул 200 с html вместо json на первой карточке: исключение вылетает из цикла,
    вторая рабочая карточка не читается, место деятельности не находится."""
    chz = _chz_with(
        [{"number": "a"}, {"number": "b"}],
        {"a": (200, b"<html>bad gateway</html>"),
         "b": (200, {"senderMod": {"fiasId": "GOOD-FIAS"}})},
    )
    assert await chz.last_mod_fias("lp") == "GOOD-FIAS"


async def test_otklonennyj_dokument_ne_dolzhen_davat_mesto_deyatelnosti():
    """Берется первый документ с fiasId, статус документа не смотрится: отклоненный ЧЗ
    вывод с ошибочным адресом запоминается за человеком навсегда."""
    chz = _chz_with(
        [{"number": "a", "status": "CHECKED_NOT_OK"}, {"number": "b", "status": "CHECKED_OK"}],
        {"a": (200, {"senderMod": {"fiasId": "BAD-FIAS"}}),
         "b": (200, {"senderMod": {"fiasId": "GOOD-FIAS"}})},
    )
    assert await chz.last_mod_fias("lp") == "GOOD-FIAS"


def test_mesto_deyatelnosti_ne_perenositsya_na_drugoj_inn_togo_zhe_cheloveka(monkeypatch):
    """Запомненное место лежит в users, а не за ИНН. Человек со вторым ИП/ООО в новом
    прогоне получает адрес первого."""
    otvety = iter(["FIAS-OF-INN-A", "FIAS-OF-INN-B"])

    async def po_ocheredi(self, pg: str = "lp") -> str:
        return next(otvety)

    monkeypatch.setattr(chz_api.ChzClient, "last_mod_fias", po_ocheredi)
    c = _client(9404)

    def vyvod(run_id: str) -> dict:
        r = c.post(f"/step4/{run_id}/build", json={"doc_type": "output"})
        assert r.status_code == 200, r.text
        return json.loads(base64.b64decode(r.json()["product_document"]))

    def novyj_progon(inn: str) -> str:
        r = c.post("/step1", data={"source": "api"}, follow_redirects=False)
        run_id = r.headers["location"].rsplit("/", 1)[-1]
        _signin(c, run_id)
        with db.connect() as conn:
            conn.execute("UPDATE runs SET inn = ? WHERE id = ?", (inn, run_id))
            conn.commit()
        return run_id

    first = vyvod(novyj_progon("7700000001"))
    second = vyvod(novyj_progon("7700000002"))
    assert first["fias_id"] == "FIAS-OF-INN-A"
    assert second["inn"] == "7700000002"
    assert second["fias_id"] == "FIAS-OF-INN-B", second["fias_id"]
