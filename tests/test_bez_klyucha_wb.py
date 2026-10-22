"""Два режима: по файлам ключ WB не нужен, по апи нужен."""
from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

import base64

from chzkiz import chz_mock, db
from chzkiz.web import app as app_mod
from chzkiz.web.app import create_app

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GS = "\\u001d"


def _code(n: int) -> str:
    return f"0104600000000{n:03d}21ABCDE{n:04d}{GS}91EE11{GS}92ZZZZ{n:04d}"


def _save(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fin(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Тип документа", "Обоснование для оплаты", "Код маркировки", "Номер сборочного задания",
               "Вайлдберриз реализовал Товар (Пр)", "Дата продажи"])
    for op, code, order, price, day in rows:
        ws.append([op, op, code, order, price, day])
    return _save(wb)


def _arch(rows) -> bytes:
    wb = openpyxl.Workbook()
    tasks = wb.active
    tasks.title = "Сборочные задания"
    tasks.append(["№ задания", "Статус задания", "Артикул"])
    sheet = wb.create_sheet("КИЗ")
    sheet.append(["№ задания", "Стикер", "КИЗ", "Номер чека", "Стоимость", "Валюта",
                  "Номер фискального накопителя", "Дата", "Тип операции", "Признак продажи"])
    for order, code, price in rows:
        tasks.append([order, "Доставлен", "A"])
        sheet.append([order, "1:2", code, "", price, "RUB", "", None, "-", "Нет"])
    return _save(wb)


@pytest.fixture
def real(monkeypatch):
    """Боевой режим: без ключа WB тут раньше приходил отказ."""
    monkeypatch.setattr(app_mod.config, "is_mock", lambda: False)


def _client(tg_id: int) -> TestClient:
    with db.connect() as conn:
        user = db.get_or_create_user(conn, tg_id)
        token = db.create_login_token(conn, user["id"])
    c = TestClient(create_app())
    c.post(f"/l/{token}", follow_redirects=True)
    return c


def _upload(c: TestClient, files):
    return c.post("/step1", data={"source": "excel"},
                  files=[("file", (n, b, XLSX)) for n, b in files], follow_redirects=False)


def _rows(run_id: str) -> list[dict]:
    with db.connect() as conn:
        return [dict(x) for x in db.get_items(conn, run_id)]


def test_arhiv_zadaniy_chitaetsya_bez_klyucha_wb(real):
    c = _client(9501)
    r = _upload(c, [("archive.xlsx", _arch([(7001, _code(11), "1 000,00")]))])
    assert r.status_code == 302, r.text[:400]
    run_id = r.headers["location"].rsplit("/", 1)[-1]
    rows = _rows(run_id)
    assert len(rows) == 1, rows
    assert rows[0]["istochnik"] == "архив", rows[0]


def test_stroki_bez_koda_otkladyvayutsya_a_ne_otbivayut_zagruzku(real):
    c = _client(9502)
    r = _upload(c, [("report.xlsx", _fin([
        ("Продажа", _code(12), 7002, "500", "2026-10-01"),
        ("Продажа", "", 7003, "600", "2026-10-02"),
    ]))])
    assert r.status_code == 302, r.text[:400]
    run_id = r.headers["location"].rsplit("/", 1)[-1]
    rows = _rows(run_id)
    assert len(rows) == 1, [(x["cis_raw"], x["istochnik"]) for x in rows]
    with db.connect() as conn:
        run = db.get_run(conn, run_id, db.get_or_create_user(conn, 9502)["id"])
    assert run["meta_failed_count"] == 1, run["meta_failed_count"]


def test_otchet_tolko_iz_fbs_bez_klyucha_daet_ponyatnuyu_oshibku(real):
    c = _client(9503)
    r = _upload(c, [("report.xlsx", _fin([("Продажа", "", 7004, "500", "2026-10-03")]))])
    assert r.status_code == 400, r.status_code
    assert "нужен ключ WB" in r.text, r.text[:400]


def test_shag3_preduprezhdaet_chto_statusy_wb_ne_podtyanuty(monkeypatch):
    """Загрузка идет в боевом режиме без ключа, сверка потом в моке."""
    c = _client(9504)
    monkeypatch.setattr(app_mod.config, "is_mock", lambda: False)
    r = _upload(c, [("archive.xlsx", _arch([(7005, _code(13), "1 000,00")]))])
    run_id = r.headers["location"].rsplit("/", 1)[-1]
    monkeypatch.undo()
    p = c.post(f"/step2/{run_id}/auth-key").json()
    sig = base64.b64encode(f"MOCKSIG:{p['data']}".encode()).decode("ascii")
    assert c.post(
        f"/step2/{run_id}/signin",
        json={"uuid": p["uuid"], "signature": sig, "inn": chz_mock.DEMO_OUR_INN, "owner": "Тест"},
    ).status_code == 200
    page = c.get(f"/step3/{run_id}").text
    assert "статусы заданий" in page, page[:600]
