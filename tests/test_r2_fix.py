"""Повторы архива в загрузке, предпросмотр кода без цены, год вне диапазона в периоде."""
from __future__ import annotations

import io

import openpyxl

from chzkiz import chz_api, chz_mock
from tests.test_attacks_mod import _signin
from tests.test_bez_klyucha_wb import _arch, _client, _code, _fin
from tests.test_v2_a2 import INN, _build, _norm, _run
from tests.test_v4_shag1 import _api_klient, _layout, _post


def _arh_bez_ceny():
    arh = _arch([(5001, _code(1), None), (5002, _code(2), None)])
    fin = _fin([("Продажа", _code(1), 5001, 1500, "2026-10-01"), ("Продажа", _code(2), 5002, 1800, "2026-10-02")])
    return arh, fin


def test_arhiv_v_oboih_polyah_cena_iz_finotcheta():
    arh, fin = _arh_bez_ceny()
    odin = _layout(_post(_client(9801), [("file", "arh.xlsx", arh), ("file_fin", "fin.xlsx", fin)]))
    dva = _layout(_post(_client(9802), [("file", "arh.xlsx", arh), ("file_fin", "arh.xlsx", arh), ("file_fin", "fin.xlsx", fin)]))
    assert [x[2] for x in odin] == [150000, 180000]
    assert dva == odin


def test_arhiv_v_oboih_polyah_kod_uhodit_v_vyvod():
    arh, fin = _arh_bez_ceny()
    for n in (1, 2):
        chz_mock.seed_cis(_norm(n), chz_api.STATUS_INTRODUCED, INN)
    c = _client(9803)
    run_id = _post(c, [("file", "arh.xlsx", arh), ("file_fin", "arh.xlsx", arh), ("file_fin", "fin.xlsx", fin)])
    _signin(c, run_id)
    r = _build(c, run_id, codes=[_norm(1)])
    assert r.status_code == 200, r.text


def test_peresekayushchiesya_arhivy_cena_u_obshchego_zadaniya():
    a1 = _arch([(5001, _code(1), None), (5002, _code(2), None)])
    a2 = _arch([(5002, _code(2), None), (5003, _code(3), None)])
    fin = _fin([("Продажа", _code(1), 5001, 1500, "2026-10-01"), ("Продажа", _code(2), 5002, 1800, "2026-10-02"),
                ("Продажа", _code(3), 5003, 1900, "2026-10-03")])
    lay = _layout(_post(_client(9804), [("file", "sent.xlsx", a1), ("file", "okt.xlsx", a2), ("file_fin", "fin.xlsx", fin)]))
    assert [x[2] for x in lay] == [150000, 180000, 190000]


def test_tri_kopii_arhiva_ne_dvoyat_stroki():
    arh = _arch([(5001, _code(1), 1500), (5002, _code(2), 1800)])
    fin = _fin([("Продажа", _code(1), 5001, 1500, "2026-10-01"), ("Продажа", _code(2), 5002, 1800, "2026-10-02")])
    lay = _layout(_post(_client(9805), [("file", "a.xlsx", arh), ("file", "a_kopiya.xlsx", arh), ("file_fin", "a.xlsx", arh),
                                        ("file_fin", "fin.xlsx", fin)]))
    assert len(lay) == 2, lay


def test_predprosmotr_kod_bez_ceny_v_liste_bez_izmeneniy():
    c, run_id = _run(9806, sales=[1, 2], returns=[5], prices={2: ""})
    assert _build(c, run_id, codes=[_norm(2)]).status_code == 400
    wb = openpyxl.load_workbook(io.BytesIO(c.get(f"/preview/{run_id}").content))
    budut = [row[0] for row in wb["Будут изменены"].iter_rows(min_row=2, values_only=True)]
    bez = {row[0]: row[4] for row in wb["Без изменений"].iter_rows(min_row=2, values_only=True)}
    assert _norm(2) not in budut
    assert _norm(1) in budut
    assert "нет цены" in bez[_norm(2)]


def test_god_0001_v_periode_ne_ronyaet_shag_1(monkeypatch):
    c = _api_klient(monkeypatch, 9807, [])
    r = c.post("/step1", data={"source": "api", "date_from": "0001-01-01", "date_to": "0001-01-05"}, follow_redirects=False)
    assert r.status_code in (302, 400)
