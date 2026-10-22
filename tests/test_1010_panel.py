from __future__ import annotations

import base64
import io
import json
import re
import time

import httpx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from chzkiz import chz_api, chz_errors, chz_mock, config, db, excel_out, kiz, sverka, wb_api
from chzkiz.web.app import create_app

INN = chz_mock.DEMO_OUR_INN
GTIN = "04600000000017"
CODE_RAW = f"01{GTIN}21PANEL1010001\\u001d91EE01\\u001d92" + "A" * 44
HDR = ("Код маркировки", "Тип документа", "Вайлдберриз реализовал Товар (Пр)")
REAL_REASON = "Выявлены коды идентификации с недопустимым статусом"


def _mocksig(value):
    return base64.b64encode(f"MOCKSIG:{value}".encode()).decode("ascii")


def _signin(c, run_id, inn):
    p = c.post(f"/step2/{run_id}/auth-key").json()
    return c.post(
        f"/step2/{run_id}/signin",
        json={"uuid": p["uuid"], "signature": _mocksig(p["data"]), "inn": inn, "owner": "Тест"},
    )


def _send(c, run_id, d):
    return c.post(
        f"/step4/{run_id}/send",
        json={"doc_id": d["doc_id"], "chz_type": d["chz_type"],
              "product_document": d["product_document"], "signature": _mocksig(d["product_document"])},
    )


def _excel_run(c, tg_id, rows):
    _login(c, tg_id)
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post("/step1", data={"source": "excel"},
               files={"file": ("r.xlsx", buf.getvalue(), "application/octet-stream")},
               follow_redirects=False)
    assert r.status_code == 302
    return r.headers["location"].rsplit("/", 1)[-1]


@pytest.mark.parametrize(
    "text, expected",
    [
        (REAL_REASON, "Повторно выводить не нужно"),
        ("Код не найден в системе", "криптохвостом"),
        ("Участник не является владельцем кода", "другим ИНН"),
        ("Ошибка проверки подписи", "флешку"),
        ("Не указано место осуществления деятельности", "место деятельности"),
        ("Неверная дата документа", "дат"),
        ("Не указана цена товара", "цен"),
    ],
)
def test_explain_known(text, expected):
    answer = chz_errors.explain(text)
    assert expected in answer
    assert "ё" not in answer and "«" not in answer and "—" not in answer


def test_explain_unknown_keeps_original_and_asks_screen():
    answer = chz_errors.explain("Странный сбой 42")
    assert answer.startswith("Странный сбой 42")
    assert "Пришлите скрин в поддержку" in answer


def test_explain_empty():
    assert chz_errors.explain("") == ""
    assert chz_errors.explain(None) == ""


def _login(c, tg_id):
    with db.connect() as conn:
        user = db.get_or_create_user(conn, tg_id)
        token = db.create_login_token(conn, user["id"])
    assert c.post(f"/l/{token}", follow_redirects=True).status_code == 200


def _demo_run(c, tg_id):
    _login(c, tg_id)
    r = c.post("/step1", data={"source": "api"}, follow_redirects=False)
    run_id = r.headers["location"].rsplit("/", 1)[-1]
    assert _signin(c, run_id, INN).status_code == 200
    return run_id


def test_step4_has_no_paid_question():
    c = TestClient(create_app())
    run_id = _demo_run(c, 10101)
    page = c.get(f"/step4/{run_id}").text
    assert "Покупатель оплатил" not in page
    assert 'id="paid"' not in page and 'id="primary-number"' not in page


def test_return_doc_defaults_to_unpaid():
    c = TestClient(create_app())
    run_id = _demo_run(c, 10102)
    r = c.post(f"/step4/{run_id}/build", json={"doc_type": "return", "paid": False})
    assert r.status_code == 200
    doc = json.loads(base64.b64decode(r.json()["product_document"]))
    assert doc["products_list"] and all(p["paid"] is False for p in doc["products_list"])
    assert "document_number" not in doc


def test_app_js_sends_return_unpaid_and_shows_explain():
    path = config.__file__.replace("config.py", "web/static/app.js")
    js = open(path, encoding="utf-8").read()
    assert "payload.paid = false" in js and "getElementById('paid')" not in js
    assert "Что это значит" in js


def test_rejected_code_shows_explanation_on_page_and_in_excel():
    code_norm = kiz.normalize_kiz(CODE_RAW)
    chz_mock.seed_cis(code_norm, chz_api.STATUS_INTRODUCED, INN)
    chz_mock.seed_reject(code_norm)
    c = TestClient(create_app())
    run_id = _excel_run(c, 10103, [HDR, (CODE_RAW, "Продажа", 1549)])
    assert _signin(c, run_id, INN).status_code == 200
    doc = c.post(f"/step4/{run_id}/build", json={"doc_type": "output"}).json()
    assert _send(c, run_id, doc).status_code == 200
    status = c.get(f"/step4/{run_id}/status").json()
    bad = [d for d in status["documents"] if d["errors"]]
    assert bad and "отклонен" in bad[0]["errors"]
    assert "Пришлите скрин в поддержку" in bad[0]["explain"]

    r = c.get(f"/report/{run_id}")
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Ошибки"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][-1] == "Что делать"
    assert "Пришлите скрин в поддержку" in rows[1][-1]
    assert "отклонен" in rows[1][5]


def test_excel_error_sheet_explains_real_chz_reason():
    row = {"cis_raw": "x", "cis_norm": "x", "op": "sale", "group": sverka.GROUP_OUTPUT,
           "reason": "", "chz_status": "INTRODUCED"}
    docs = [{"doc_type": chz_api.DOC_TYPE_RECEIPT, "status": chz_api.DOC_STATUS_CHECKED_NOT_OK,
             "errors": REAL_REASON, "codes_json": json.dumps(["x"])}]
    data = excel_out.build_report([row], docs)
    ws = openpyxl.load_workbook(io.BytesIO(data))["Ошибки"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][5] == REAL_REASON
    assert "Повторно выводить не нужно" in rows[1][6]


def test_session_lives_30_days():
    assert config.SESSION_TTL_HOURS == 720
    with db.connect() as conn:
        user = db.get_or_create_user(conn, 10104)
        sid = db.create_session(conn, user["id"])
        left = conn.execute("SELECT expires_at FROM sessions WHERE session_id = ?", (sid,)).fetchone()["expires_at"]
    assert left - time.time() > 29 * 86400


def _patch_wb(monkeypatch, handler):
    real = httpx.AsyncClient

    def patched(*args, **kwargs):
        if kwargs.get("transport") is None:
            kwargs["transport"] = httpx.MockTransport(handler)
        return real(*args, **kwargs)

    monkeypatch.setattr(httpx, "AsyncClient", patched)


@pytest.mark.parametrize("cancel_status", ["canceled_by_client", "declined_by_client"])
@pytest.mark.parametrize("cancel_date, sale_date", [("2026-09-20", "2026-09-22"), ("2026-09-22", "2026-09-22")])
async def test_api_cancel_then_resale_goes_to_output(monkeypatch, cancel_status, cancel_date, sale_date):
    code = "010460312345678921PANEL1010xxx"
    statusy = {301: cancel_status, 302: "sold"}
    dates = {301: cancel_date, 302: sale_date}

    def handler(request: httpx.Request) -> httpx.Response:
        p = request.url.path
        if p == "/api/v3/orders":
            return httpx.Response(200, json={"next": 0, "orders": [
                {"id": i, "price": 1000, "createdAt": dates[i] + "T10:00:00Z"} for i in statusy]})
        if p == "/api/v3/orders/status":
            return httpx.Response(200, json={"orders": [{"id": i, "wbStatus": s} for i, s in statusy.items()]})
        if p == "/api/marketplace/v3/orders/meta":
            nado = json.loads(request.content)["orders"]
            return httpx.Response(200, json={"orders": [
                {"id": i, "meta": {"sgtin": {"value": [code]}}} for i in nado]})
        return httpx.Response(404, json={})

    _patch_wb(monkeypatch, handler)
    items = await wb_api.WbClient("test-key-not-real").fetch_items(date_from=wb_api.msk_day_unix("2026-09-01"), date_to=wb_api.msk_day_unix("2026-10-09", end=True))
    assert wb_api.OP_SALE in [it["op"] for it in items]
    norm = kiz.normalize_kiz(code)
    info = {norm: {"cis": norm, "status": chz_api.STATUS_INTRODUCED, "ownerInn": INN}}
    rows = sverka.classify_items(items, info, INN)
    assert [r["group"] for r in rows] == [sverka.GROUP_OUTPUT], rows
