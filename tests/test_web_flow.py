"""Сквозной прогон в демо-режиме (CHZ_MODE=mock): ссылка входа -> сверка -> вывод -> возврат -> отчет."""
from __future__ import annotations

import base64
import io
import json
from datetime import date, datetime, timedelta

import httpx
import openpyxl
from fastapi.testclient import TestClient

from chzkiz import chz_api, chz_mock, db, kiz, sverka, wb_api
from chzkiz.web.app import create_app

client = TestClient(create_app())


def _mocksig(value: str) -> str:
    return base64.b64encode(f"MOCKSIG:{value}".encode()).decode("ascii")


def _login_link_for(tg_id: int) -> str:
    with db.connect() as conn:
        user = db.get_or_create_user(conn, tg_id)
        token = db.create_login_token(conn, user["id"])
    return f"/l/{token}"


def test_health():
    r = client.get("/health")
    assert r.status_code == 200
    assert r.json() == {"ok": True}


def test_step1_requires_session():
    r = client.get("/step1", follow_redirects=False)
    assert r.status_code == 401


def test_full_demo_flow_output_and_return():
    link = _login_link_for(9001)
    r = client.post(link, follow_redirects=True)
    assert r.status_code == 200
    assert r.url.path == "/step1"

    r = client.post("/step1", data={"source": "api"}, follow_redirects=False)
    assert r.status_code == 302
    run_id = r.headers["location"].split("/")[-1]

    r = client.post(f"/step2/{run_id}/auth-key")
    assert r.status_code == 200
    payload = r.json()
    signature = _mocksig(payload["data"])
    r = client.post(
        f"/step2/{run_id}/signin",
        json={
            "uuid": payload["uuid"],
            "signature": signature,
            "inn": chz_mock.DEMO_OUR_INN,
            "owner": "Тест Тестович",
        },
    )
    assert r.status_code == 200

    r = client.get(f"/step3/{run_id}")
    assert r.status_code == 200
    assert "Вывести" in r.text

    r = client.get(f"/step4/{run_id}")
    assert r.status_code == 200

    r = client.post(f"/step4/{run_id}/build", json={"doc_type": "output"})
    assert r.status_code == 200
    out_doc = r.json()
    sig = _mocksig(out_doc["product_document"])
    r = client.post(
        f"/step4/{run_id}/send",
        json={
            "doc_id": out_doc["doc_id"],
            "chz_type": out_doc["chz_type"],
            "product_document": out_doc["product_document"],
            "signature": sig,
        },
    )
    assert r.status_code == 200

    r = client.post(f"/step4/{run_id}/build", json={"doc_type": "return", "paid": False})
    assert r.status_code == 200
    ret_doc = r.json()
    sig2 = _mocksig(ret_doc["product_document"])
    r = client.post(
        f"/step4/{run_id}/send",
        json={
            "doc_id": ret_doc["doc_id"],
            "chz_type": ret_doc["chz_type"],
            "product_document": ret_doc["product_document"],
            "signature": sig2,
        },
    )
    assert r.status_code == 200

    r = client.get(f"/step4/{run_id}/status")
    assert r.status_code == 200
    body = r.json()
    assert body["done"] is True

    r = client.get(f"/report/{run_id}")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/")

    wb_report = openpyxl.load_workbook(io.BytesIO(r.content))

    def sheet_codes(name: str) -> list[str]:
        ws = wb_report[name]
        return [row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]

    output_code = kiz.normalize_kiz(wb_api._fake_cis("781001"))
    return_code = kiz.normalize_kiz(wb_api._fake_cis("781003"))
    assert sheet_codes("Выведено") == [output_code], "лист Выведено не совпадает с принятым ЧЗ документом"
    assert sheet_codes("Возвращено") == [return_code], "лист Возвращено не совпадает с принятым ЧЗ документом"
    error_codes = sheet_codes("Ошибки")
    assert return_code not in error_codes, "отправленный возврат не должен также висеть в Ошибках"
    assert len(error_codes) == 5, f"демо-набор дает 5 ошибочных строк, в отчете {len(error_codes)}"


def test_run_not_visible_to_other_session():
    link = _login_link_for(9002)
    client.post(link, follow_redirects=True)
    r = client.post("/step1", data={"source": "api"}, follow_redirects=False)
    run_id = r.headers["location"].split("/")[-1]

    other = TestClient(create_app())
    other_link = _login_link_for(9003)
    other.post(other_link, follow_redirects=True)
    r = other.get(f"/step3/{run_id}")
    assert r.status_code == 404


def test_rejected_doc_goes_to_error_sheet_not_output():
    """Отказ ЧЗ (CHECKED_NOT_OK) на реально отправленный документ - код должен лечь на
    лист Ошибки, а не остаться в Выведено: исход решает статус документа, не намерение."""
    # текстовый эскейп GS ("\u001d" из 6 знаков), не байт - иначе openpyxl не пишет ячейку.
    code_raw = "01" + wb_api._FAKE_GTIN + "21781900\\u001d91EE01\\u001d92AbCdEf=="
    code_norm = kiz.normalize_kiz(code_raw)
    chz_mock.seed_cis(code_norm, chz_api.STATUS_INTRODUCED, chz_mock.DEMO_OUR_INN)
    chz_mock.seed_reject(code_norm)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Код маркировки", "Тип документа", "Вайлдберриз реализовал Товар (Пр)"])
    ws.append([code_raw, "Продажа", 1549])
    buf = io.BytesIO()
    wb.save(buf)

    c = TestClient(create_app())
    link = _login_link_for(9006)
    c.post(link, follow_redirects=True)
    r = c.post(
        "/step1",
        data={"source": "excel"},
        files={"file": ("report.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    run_id = r.headers["location"].split("/")[-1]

    r = c.post(f"/step2/{run_id}/auth-key")
    payload = r.json()
    r = c.post(
        f"/step2/{run_id}/signin",
        json={
            "uuid": payload["uuid"],
            "signature": _mocksig(payload["data"]),
            "inn": chz_mock.DEMO_OUR_INN,
            "owner": "Тест",
        },
    )
    assert r.status_code == 200

    r = c.get(f"/step3/{run_id}")
    assert "Вывести" in r.text

    r = c.post(f"/step4/{run_id}/build", json={"doc_type": "output"})
    assert r.status_code == 200
    out_doc = r.json()
    r = c.post(
        f"/step4/{run_id}/send",
        json={
            "doc_id": out_doc["doc_id"],
            "chz_type": out_doc["chz_type"],
            "product_document": out_doc["product_document"],
            "signature": _mocksig(out_doc["product_document"]),
        },
    )
    assert r.status_code == 200

    r = c.get(f"/step4/{run_id}/status")
    assert r.status_code == 200
    body = r.json()
    assert body["done"] is True
    assert body["documents"][0]["status"] == chz_api.DOC_STATUS_CHECKED_NOT_OK

    r = c.get(f"/report/{run_id}")
    assert r.status_code == 200
    wb_report = openpyxl.load_workbook(io.BytesIO(r.content))

    def sheet_codes(name: str) -> list[str]:
        ws = wb_report[name]
        return [row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]

    assert code_norm in sheet_codes("Ошибки"), "отклоненный ЧЗ код должен лежать в Ошибках"
    assert code_norm not in sheet_codes("Выведено"), "отклоненный ЧЗ код не должен висеть в Выведено"


def test_excel_upload_flow_creates_run():
    from io import BytesIO

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Код маркировки", "Тип документа"])
    ws.append(["01460312345678902121serial\\u001d91EE01\\u001d92sig==", "Продажа"])
    buf = BytesIO()
    wb.save(buf)

    link = _login_link_for(9004)
    upload_client = TestClient(create_app())
    upload_client.post(link, follow_redirects=True)
    r = upload_client.post(
        "/step1",
        data={"source": "excel"},
        files={"file": ("report.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    assert r.status_code == 302
    assert "/step2/" in r.headers["location"]


def test_login_link_get_shows_button_and_keeps_token():
    link = _login_link_for(9012)
    c = TestClient(create_app())
    r = c.get(link, follow_redirects=False)
    assert r.status_code == 200
    assert f'action="{link}"' in r.text
    r = c.get(link, follow_redirects=False)
    assert r.status_code == 200, "повторный GET (предпросмотр ссылки) не должен гасить ссылку"
    r = c.post(link, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/step1"
    assert c.get("/step1").status_code == 200


def test_login_link_used_but_session_alive_goes_to_step1():
    link = _login_link_for(9013)
    c = TestClient(create_app())
    assert c.post(link, follow_redirects=False).status_code == 303
    r = c.get(link, follow_redirects=False)
    assert r.status_code == 302 and r.headers["location"] == "/step1"
    r = c.post(link, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/step1"
    stranger = TestClient(create_app())
    assert stranger.get(link, follow_redirects=False).status_code == 400
    assert stranger.post(link, follow_redirects=False).status_code == 400


def test_excel_fbs_report_takes_codes_by_order_and_drops_fbo(monkeypatch):
    from io import BytesIO

    from chzkiz.web import app as app_mod

    seen: list[dict] = []

    async def fake_fill(wb_key_enc, items):
        seen.extend(items)
        return [{**it, "cis_raw": "0104603123456789021fbs1", "need_meta": False} for it in items], 0

    monkeypatch.setattr(app_mod, "_fill_fbs_codes", fake_fill)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Код маркировки", "Тип документа", "Обоснование для оплаты", "Номер сборочного задания"])
    ws.append(["01460312345678902121fbo\\u001d91EE01\\u001d92sig==", "Продажа", "Продажа", ""])
    ws.append(["", "Продажа", "Продажа", 5812549572])
    ws.append(["", "Продажа", "Возмещение за выдачу и возврат товаров на ПВЗ", 5812549573])
    buf = BytesIO()
    wb.save(buf)

    upload_client = TestClient(create_app())
    upload_client.post(_login_link_for(9005), follow_redirects=True)
    r = upload_client.post(
        "/step1",
        data={"source": "excel"},
        files={"file": ("report.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    assert r.status_code == 302 and "/step2/" in r.headers["location"]
    assert [it["order_id"] for it in seen] == [5812549572]


def test_excel_report_with_only_fbo_rows_says_nothing_to_check():
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Код маркировки", "Тип документа", "Обоснование для оплаты", "Номер сборочного задания"])
    ws.append(["01460312345678902121fbo\\u001d91EE01\\u001d92sig==", "Продажа", "Продажа", ""])
    buf = BytesIO()
    wb.save(buf)

    upload_client = TestClient(create_app())
    upload_client.post(_login_link_for(9006), follow_redirects=True)
    r = upload_client.post(
        "/step1",
        data={"source": "excel"},
        files={"file": ("report.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    assert "выводит сам WB" in r.text


def test_return_paid_needs_primary_document():
    c = TestClient(create_app())
    c.post(_login_link_for(9010), follow_redirects=True)
    r = c.post("/step1", data={"source": "api"}, follow_redirects=False)
    run_id = r.headers["location"].split("/")[-1]
    payload = c.post(f"/step2/{run_id}/auth-key").json()
    r = c.post(
        f"/step2/{run_id}/signin",
        json={
            "uuid": payload["uuid"],
            "signature": _mocksig(payload["data"]),
            "inn": chz_mock.DEMO_OUR_INN,
            "owner": "Тест",
        },
    )
    assert r.status_code == 200
    assert "Цена, р." in c.get(f"/step3/{run_id}").text
    assert 'id="paid"' not in c.get(f"/step4/{run_id}").text

    for paid in (None, "true", 1):
        body = {"doc_type": "return"} if paid is None else {"doc_type": "return", "paid": paid}
        assert c.post(f"/step4/{run_id}/build", json=body).status_code == 400, paid
    assert c.post(f"/step4/{run_id}/build", json={"doc_type": "foo"}).status_code == 400

    r = c.post(f"/step4/{run_id}/build", json={"doc_type": "return", "paid": False})
    assert r.status_code == 200
    doc = json.loads(base64.b64decode(r.json()["product_document"]))
    assert all(p == {"ki": p["ki"], "paid": False} for p in doc["products_list"])

    full = {
        "doc_type": "return",
        "paid": True,
        "primary_type": "OTHER",
        "primary_number": "862067877",
        "primary_date": date.today().isoformat(),
        "primary_name": "Отчет о реализации Wildberries",
    }
    for bad in (
        {"primary_number": ""},
        {"primary_number": "7" * 256},
        {"primary_date": (date.today() + timedelta(days=1)).isoformat()},
        {"primary_date": "не дата"},
        {"primary_type": "BOGUS"},
        {"primary_name": ""},
        {"primary_date": (_five_years_ago() - timedelta(days=1)).isoformat()},
        {"primary_number": "​"},
        {"primary_number": "⁠­"},
        {"primary_name": "­"},
    ):
        assert c.post(f"/step4/{run_id}/build", json={**full, **bad}).status_code == 400, bad
    for good in ({"primary_number": "7" * 255}, {"primary_date": _five_years_ago().isoformat()}):
        assert c.post(f"/step4/{run_id}/build", json={**full, **good}).status_code == 200, good

    r = c.post(f"/step4/{run_id}/build", json={**full, "primary_date": date.today().strftime("%Y%m%d")})
    assert r.status_code == 200
    doc = json.loads(base64.b64decode(r.json()["product_document"]))
    assert doc["return_type"] == "REMOTE_SALE_RETURN"
    item = doc["products_list"][0]
    assert item["paid"] is True
    assert item["primary_document_type"] == "OTHER"
    assert item["primary_document_number"] == "862067877"
    assert item["primary_document_date"] == date.today().isoformat()
    assert item["primary_document_custom_name"] == "Отчет о реализации Wildberries"


GTIN = "04600000000017"
TAIL = "\\u001d91EE01\\u001d92" + "A" * 44


def _five_years_ago() -> date:
    today = date.today()
    try:
        return today.replace(year=today.year - 5)
    except ValueError:
        return today.replace(year=today.year - 5, day=28)


def _signin(c: TestClient, run_id: str) -> None:
    payload = c.post(f"/step2/{run_id}/auth-key").json()
    r = c.post(
        f"/step2/{run_id}/signin",
        json={
            "uuid": payload["uuid"],
            "signature": _mocksig(payload["data"]),
            "inn": chz_mock.DEMO_OUR_INN,
            "owner": "Тест",
        },
    )
    assert r.status_code == 200, r.text[:200]


def _excel_run(tg_id: int, rows: list[tuple]) -> tuple[TestClient, str]:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    c = TestClient(create_app())
    c.post(_login_link_for(tg_id), follow_redirects=True)
    r = c.post(
        "/step1",
        data={"source": "excel"},
        files={"file": ("report.xlsx", buf.getvalue(), "application/octet-stream")},
        follow_redirects=False,
    )
    assert r.status_code == 302, r.text[:200]
    run_id = r.headers["location"].split("/")[-1]
    _signin(c, run_id)
    return c, run_id


def test_step3_shows_report_date_and_price_in_any_date_format():
    codes = [f"01{GTIN}21WFLOW3S00000{i}" for i in range(1, 5)]
    for code in codes[:3]:
        chz_mock.seed_cis(code, chz_api.STATUS_INTRODUCED, chz_mock.DEMO_OUR_INN)
    chz_mock.seed_cis(codes[3], chz_api.STATUS_RETIRED, chz_mock.DEMO_OUR_INN)
    rows = [
        ("Код маркировки", "Тип документа", "Дата продажи", "Вайлдберриз реализовал Товар (Пр)"),
        (codes[0] + TAIL, "Продажа", datetime(2026, 10, 5), 1549),
        (codes[1] + TAIL, "Продажа", "05.10.2026", 1549),
        (codes[2] + TAIL, "Продажа", "2026-10-05", 1549),
        (codes[3] + TAIL, "Продажа", "28.09.2026", 1549),
        (codes[3] + TAIL, "Возврат", "05.10.2026", 1549),
    ]
    c, run_id = _excel_run(9020, rows)
    with db.connect() as conn:
        items = [dict(r) for r in db.get_items(conn, run_id)]
    assert {r["date"] for r in items if r["cis_norm"] in codes[:3]} == {"2026-10-05"}
    assert sverka.GROUP_RETURN in [r["group_name"] for r in items if r["cis_norm"] == codes[3]]
    html = c.get(f"/step3/{run_id}").text
    assert "05.10.2026" in html and "10.05.2026" not in html
    assert "1549.00" in html and "154900" not in html


def test_output_without_price_is_not_built():
    code = f"01{GTIN}21WFLOW3S000009"
    chz_mock.seed_cis(code, chz_api.STATUS_INTRODUCED, chz_mock.DEMO_OUR_INN)
    c, run_id = _excel_run(9021, [("Код маркировки", "Тип документа", "Дата продажи"), (code + TAIL, "Продажа", "2026-10-05")])
    r = c.post(f"/step4/{run_id}/build", json={"doc_type": "output"})
    assert r.status_code == 400
    assert code in r.json()["detail"]


def test_send_takes_plain_text_doc_id_from_chz(monkeypatch):
    orig = httpx.ASGITransport.handle_async_request

    async def plain_id(self, request):
        resp = await orig(self, request)
        if request.method == "POST" and request.url.path.endswith(chz_api.DOC_CREATE_PATH) and resp.status_code == 200:
            await resp.aread()
            return httpx.Response(
                200, text=resp.json()["id"], headers={"content-type": "text/plain;charset=UTF-8"}, request=request
            )
        return resp

    monkeypatch.setattr(httpx.ASGITransport, "handle_async_request", plain_id)
    c = TestClient(create_app())
    c.post(_login_link_for(9022), follow_redirects=True)
    r = c.post("/step1", data={"source": "api"}, follow_redirects=False)
    run_id = r.headers["location"].split("/")[-1]
    _signin(c, run_id)
    b = c.post(f"/step4/{run_id}/build", json={"doc_type": "output"}).json()
    send = {
        "doc_id": b["doc_id"],
        "chz_type": b["chz_type"],
        "product_document": b["product_document"],
        "signature": _mocksig(b["product_document"]),
    }
    first = c.post(f"/step4/{run_id}/send", json=send)
    assert first.status_code == 200, first.text[:200]
    chz_id = first.json()["chz_doc_id"]
    assert chz_id in chz_mock._docs
    assert c.post(f"/step4/{run_id}/send", json=send).json()["chz_doc_id"] == chz_id
    assert len(chz_mock._docs) == 1
