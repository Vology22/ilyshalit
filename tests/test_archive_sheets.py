"""Файл выгрузки архива сборочных заданий WB: коды на втором листе "КИЗ", операция по статусам WB."""
from __future__ import annotations

import asyncio
from io import BytesIO

import openpyxl

from chzkiz import excel_in, sverka, wb_api
from chzkiz.web import app as app_mod

GS = "\\u001d"


def _code(n: int) -> str:
    return f"0104600000000{n:03d}21ABCDE{n:04d}{GS}91EE11{GS}92ZZZZ{n:04d}"


def make_archive(n: int = 5, empty_last: bool = False) -> bytes:
    wb = openpyxl.Workbook()
    tasks = wb.active
    tasks.title = "Сборочные задания"
    tasks.append(["№ задания", "Статус задания", "Артикул"])
    kiz = wb.create_sheet("КИЗ")
    kiz.append(["№ задания", "Стикер", "КИЗ", "Номер чека", "Стоимость", "Валюта",
                "Номер фискального накопителя", "Дата", "Тип операции", "Признак продажи юрлицу"])
    for i in range(n):
        order = 1000 + i
        tasks.append([order, "Доставлен", "A"])
        code = "" if empty_last and i == n - 1 else _code(i)
        kiz.append([order, "123:456", code, "", "1 234,50", "RUB", "", None, "-", "Нет"])
    return _save(wb)


def _save(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_finreport() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Тип документа", "Обоснование для оплаты", "Код маркировки", "Номер сборочного задания",
               "Вайлдберриз реализовал Товар (Пр)", "Дата продажи"])
    ws.append(["Продажа", "Продажа", _code(1), 1001, "500", "2026-10-01"])
    return _save(wb)


def test_two_sheet_archive_gives_all_codes():
    items = excel_in.parse_report(make_archive(5))
    assert len(items) == 5
    assert all(it["cis_raw"] and it["archive"] and it["op"] == wb_api.OP_SALE for it in items)
    assert items[0]["order_id"] == 1000
    assert items[0]["product_cost"] == 123450
    assert all(it["date"] == "" and not it["fbo"] for it in items)


def test_archive_row_without_code_is_skipped():
    assert len(excel_in.parse_report(make_archive(5, empty_last=True))) == 4


def test_finreport_parsed_as_before():
    items = excel_in.parse_report(make_finreport())
    assert len(items) == 1
    it = items[0]
    assert it["op"] == wb_api.OP_SALE and it["order_id"] == 1001
    assert it["date"] == "2026-10-01" and it["product_cost"] == 50000
    assert "archive" not in it


def test_no_kiz_anywhere_still_errors():
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.create_sheet("x").append(["c", "d"])
    try:
        excel_in.parse_report(_save(wb))
    except excel_in.ExcelColumnError as e:
        assert e.no_kiz
    else:
        raise AssertionError("ожидалась ошибка")


class _FakeClient:
    def __init__(self, key):
        pass

    async def get_orders_status(self, ids):
        table = {1000: "sold", 1001: "canceled_by_client", 1002: "declined_by_client",
                 1003: "ready_for_pickup"}
        return {i: {"wbStatus": table[i]} for i in ids if i in table}


def _statuses(monkeypatch, items):
    monkeypatch.setattr(app_mod.wb_api, "WbClient", _FakeClient)
    monkeypatch.setattr(app_mod, "_decrypt", lambda v: v)
    return asyncio.run(app_mod._apply_wb_statuses("k", items))


def test_wb_statuses_set_operation(monkeypatch):
    items = _statuses(monkeypatch, excel_in.parse_report(make_archive(5)))
    by_order = {it["order_id"]: it for it in items}
    assert by_order[1000]["op"] == wb_api.OP_SALE and not by_order[1000].get("reason")
    assert by_order[1001]["op"] == wb_api.OP_RETURN
    assert by_order[1002]["op"] == wb_api.OP_RETURN
    assert by_order[1003]["reason"] == "not_sold_yet"
    assert by_order[1004]["reason"] == "not_sold_yet"  # WB статус не вернул


def test_not_sold_yet_goes_to_ok_group(monkeypatch):
    items = _statuses(monkeypatch, excel_in.parse_report(make_archive(5)))
    chz_info = {sverka.kiz.normalize_kiz(it["cis_raw"]): {"status": "INTRODUCED", "ownerInn": "1"} for it in items}
    rows = {r["order_id"]: r for r in sverka.classify_items(items, chz_info, "1")}
    assert rows[1003]["group"] == sverka.GROUP_OK and rows[1003]["reason"] == "not_sold_yet"
    assert rows[1000]["group"] == sverka.GROUP_OUTPUT
    assert rows[1001]["group"] == sverka.GROUP_OK and rows[1001]["reason"] == "otkaz_v_oborote"
    assert sverka.REASON_LABELS["not_sold_yet"] == "еще не продан, выводить рано"

def test_pri_arhive_kody_tolko_iz_arhiva():
    """Загружены оба файла: коды считаются только по архиву, финотчет своих не добавляет."""
    from fastapi.testclient import TestClient
    from chzkiz import db, excel_in, kiz
    from chzkiz.web.app import create_app

    with db.connect() as conn:
        user = db.get_or_create_user(conn, 9102)
        token = db.create_login_token(conn, user["id"])
    client = TestClient(create_app())
    client.post(f"/l/{token}", follow_redirects=True)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    arh_bytes = make_archive(5)
    fin_bytes = make_finreport()
    r = client.post(
        "/step1",
        data={"source": "excel"},
        files=[("file", ("report.xlsx", fin_bytes, xlsx)),
               ("file", ("report2.xlsx", _finreport_2_stroki(), xlsx)),
               ("file", ("archive.xlsx", arh_bytes, xlsx))],
        follow_redirects=False,
    )
    assert r.status_code == 302, r.text[:400]
    run_id = r.headers["location"].rsplit("/", 1)[-1]
    with db.connect() as conn:
        rows = db.get_items(conn, run_id)

    arh_cis = {kiz.normalize_kiz(it["cis_raw"]) for it in excel_in.parse_report(arh_bytes)}
    fin_cis = {kiz.normalize_kiz(it["cis_raw"])
               for b in (fin_bytes, _finreport_2_stroki()) for it in excel_in.parse_report(b)}
    bylo = {kiz.normalize_kiz(row["cis_raw"]) for row in rows}
    assert fin_cis - arh_cis, "в примере нужны коды отчета, которых нет в архиве"
    assert bylo == arh_cis, (len(bylo), len(arh_cis))
    # ни один код не попал в сверку дважды
    assert len(rows) == len(bylo)
    # общий код достается архиву, дату строке отдает последнее событие пары из финотчета
    obshie = arh_cis & fin_cis
    fin_dates = {kiz.normalize_kiz(it["cis_raw"]): it["date"] for it in excel_in.parse_report(fin_bytes)}
    for row in rows:
        if kiz.normalize_kiz(row["cis_raw"]) in obshie:
            assert row["date"] == fin_dates.get(kiz.normalize_kiz(row["cis_raw"]), row["date"]), row["date"]


def _finreport_2_stroki() -> bytes:
    """Финотчет на две строки, коды свои: с архивом пересечения быть не должно."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Тип документа", "Обоснование для оплаты", "Код маркировки", "Номер сборочного задания",
               "Вайлдберриз реализовал Товар (Пр)", "Дата продажи"])
    ws.append(["Продажа", "Продажа", _code(50), 7770, "500", "2026-10-01"])
    ws.append(["Продажа", "Продажа", _code(51), 7771, "600", "2026-10-02"])
    return _save(wb)


def test_shag3_schitaet_po_arhivu():
    """Загружены оба файла: на шаге 3 всего кодов столько, сколько в архиве."""
    import base64

    from fastapi.testclient import TestClient
    from chzkiz import chz_mock, db
    from chzkiz.web.app import create_app

    with db.connect() as conn:
        user = db.get_or_create_user(conn, 9103)
        token = db.create_login_token(conn, user["id"])
    client = TestClient(create_app())
    client.post(f"/l/{token}", follow_redirects=True)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = client.post(
        "/step1",
        data={"source": "excel"},
        files=[("file", ("report.xlsx", _finreport_2_stroki(), xlsx)),
               ("file", ("archive.xlsx", make_archive(1), xlsx))],
        follow_redirects=False,
    )
    assert r.status_code == 302, r.text[:400]
    run_id = r.headers["location"].rsplit("/", 1)[-1]
    p = client.post(f"/step2/{run_id}/auth-key").json()
    sig = base64.b64encode(f"MOCKSIG:{p['data']}".encode()).decode("ascii")
    assert client.post(
        f"/step2/{run_id}/signin",
        json={"uuid": p["uuid"], "signature": sig, "inn": chz_mock.DEMO_OUR_INN, "owner": "Тест"},
    ).status_code == 200
    with db.connect() as conn:
        rows = db.get_items(conn, run_id)
    arh = sum(1 for row in rows if row["istochnik"] == "архив")
    fin = sum(1 for row in rows if row["istochnik"] == "отчет")
    assert (arh, fin) == (1, 0), (arh, fin, len(rows))
    page = client.get(f"/step3/{run_id}").text
    assert "Всего кодов: 1." in page
